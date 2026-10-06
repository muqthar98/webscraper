import re
import time
from pathlib import Path
from urllib.parse import quote_plus

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError


FIELDS = [
    "Business Name", "Category", "Address", "Phone", "Website",
    "Rating", "Review Count", "Google Maps URL", "Search Query", "Search Location"
]


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def text_from_first(page, selectors):
    for selector in selectors:
        try:
            loc = page.locator(selector).first
            if loc.count() and loc.is_visible():
                value = clean(loc.inner_text(timeout=1500))
                if value:
                    return value
        except Exception:
            pass
    return ""


def attr_from_first(page, selectors, attr):
    for selector in selectors:
        try:
            loc = page.locator(selector).first
            if loc.count():
                value = clean(loc.get_attribute(attr, timeout=1500))
                if value:
                    return value
        except Exception:
            pass
    return ""


def parse_rating_reviews(page):
    rating = ""
    reviews = ""

    try:
        candidates = page.locator('[aria-label*="stars"], [aria-label*="star"]')
        for i in range(min(candidates.count(), 10)):
            label = clean(candidates.nth(i).get_attribute("aria-label"))
            m = re.search(r"([0-5](?:\.\d+)?)", label)
            if m:
                rating = m.group(1)
                break
    except Exception:
        pass

    try:
        body = clean(page.locator("body").inner_text(timeout=3000))
        m = re.search(r"([0-5](?:\.\d+)?)\s*\(?([\d,.\s]+)\)?", body)
        if m:
            rating = rating or m.group(1)
            reviews = re.sub(r"[^\d]", "", m.group(2))
    except Exception:
        pass

    if not reviews:
        try:
            review_loc = page.locator('[aria-label*="review"]').first
            label = clean(review_loc.get_attribute("aria-label"))
            m = re.search(r"([\d,]+)", label)
            if m:
                reviews = m.group(1).replace(",", "")
        except Exception:
            pass

    return rating, reviews


def extract_business(page, search_query, search_location):
    name = text_from_first(page, [
        "h1.DUwDvf",
        'h1[class*="fontHeadlineLarge"]',
        "h1",
    ])

    if not name:
        return None

    category = text_from_first(page, [
        'button[jsaction*="category"]',
        'div[class*="fontBodyMedium"] button',
    ])

    address = attr_from_first(page, [
        'button[data-item-id="address"]',
        'button[aria-label^="Address"]',
    ], "aria-label")
    address = re.sub(r"^Address:\s*", "", address, flags=re.I)

    phone = attr_from_first(page, [
        'button[data-item-id^="phone:tel:"]',
        'button[aria-label^="Phone:"]',
    ], "aria-label")
    phone = re.sub(r"^Phone:\s*", "", phone, flags=re.I)

    website = attr_from_first(page, [
        'a[data-item-id="authority"]',
        'a[aria-label^="Website"]',
    ], "href")

    maps_url = page.url
    rating, reviews = parse_rating_reviews(page)

    return {
        "Business Name": name,
        "Category": category,
        "Address": address,
        "Phone": phone,
        "Website": website,
        "Rating": rating,
        "Review Count": reviews,
        "Google Maps URL": maps_url,
        "Search Query": search_query,
        "Search Location": search_location,
    }


def collect_result_urls(page, max_results, scroll_pause):
    urls = []
    seen = set()
    panel = page.locator('div[role="feed"]').first

    for _ in range(100):
        links = page.locator('a[href*="/maps/place/"]')
        count = links.count()

        for i in range(count):
            try:
                href = links.nth(i).get_attribute("href")
                if href and "/maps/place/" in href and href not in seen:
                    seen.add(href)
                    urls.append(href)
                    if len(urls) >= max_results:
                        return urls
            except Exception:
                continue

        try:
            if panel.count():
                panel.evaluate("(el) => el.scrollTop = el.scrollHeight")
            else:
                page.mouse.wheel(0, 5000)
        except Exception:
            page.mouse.wheel(0, 5000)

        time.sleep(scroll_pause)

        try:
            body = clean(page.locator("body").inner_text(timeout=1500))
            if "You've reached the end of the list" in body:
                break
        except Exception:
            pass

    return urls


def scrape_search(context, query, location, max_results, scroll_pause):
    page = context.new_page()
    search_url = (
        "https://www.google.com/maps/search/"
        + quote_plus(f"{query} {location}")
    )

    print(f"\nSearching: {query} | {location}")
    page.goto(search_url, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(3000)

    urls = collect_result_urls(page, max_results, scroll_pause)
    print(f"Found {len(urls)} result links")

    results = []
    for index, url in enumerate(urls, 1):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(1200)
            item = extract_business(page, query, location)
            if item:
                results.append(item)
                print(f"[{index}/{len(urls)}] {item['Business Name']}")
        except PlaywrightTimeoutError:
            print(f"[{index}/{len(urls)}] timeout - skipped")
        except Exception as exc:
            print(f"[{index}/{len(urls)}] skipped: {exc}")

    page.close()
    return results


def format_excel(output_path, rows):
    df = pd.DataFrame(rows, columns=FIELDS)

    for col in FIELDS:
        df[col] = df[col].map(clean)

    df = df.drop_duplicates(subset=["Google Maps URL"], keep="first")
    df = df[df["Business Name"].ne("")].copy()

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Businesses")

    wb = load_workbook(output_path)
    ws = wb["Businesses"]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 32, "B": 24, "C": 45, "D": 20, "E": 40,
        "F": 10, "G": 14, "H": 60, "I": 24, "J": 24
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if ws.max_row >= 2:
        table = Table(displayName="BusinessData", ref=ws.dimensions)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Google Maps Business Scraper"
    summary["A1"].font = Font(bold=True, size=16)
    summary["A3"] = "Total businesses"
    summary["B3"] = len(df)
    summary["A4"] = "Businesses with phone"
    summary["B4"] = int(df["Phone"].ne("").sum())
    summary["A5"] = "Businesses with website"
    summary["B5"] = int(df["Website"].ne("").sum())

    wb.save(output_path)
    return len(df)
